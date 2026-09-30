from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.oxml.ns import qn

p = Path(r"D:\SmartF\output\word\SmartFoodRescue_IEEE_Research_Paper.docx")
d = Document(p)
text = "\n".join(x.text for x in d.paragraphs)
text += "\n" + "\n".join(c.text for t in d.tables for row in t.rows for c in row.cells)

required = [
    "Divya Ambre", "divyaambre21@gmail.com",
    "Sumit Ghogare", "sumitghogare2007@gmail.com",
    "Tanvi Patil", "tanvip1117@gmail.com",
    "Tanvi Shrirame", "tanavishrirame857@gmail.com",
    "I INTRODUCTION", "X REFERENCES", "DISTRIBUTED",
]
forbidden = ["Author Name(s)", "author.email@example.com", "Department and Institution Name"]

with ZipFile(p) as z:
    xml = z.read("word/document.xml").decode("utf-8")

all_paragraphs = list(d.paragraphs)
for table in d.tables:
    for row in table.rows:
        for cell in row.cells:
            all_paragraphs.extend(cell.paragraphs)

colored_runs = []
underlined_runs = []
for paragraph in all_paragraphs:
    for run in paragraph.runs:
        color = run.font.color.rgb
        if color is not None and str(color) != "000000":
            colored_runs.append((run.text, str(color)))
        if run.font.underline:
            underlined_runs.append(run.text)

print("required=", {x: x in text for x in required})
print("forbidden=", {x: x in text for x in forbidden})
print("sections=", len(d.sections))
print("tables=", len(d.tables))
print("inline_shapes=", len(d.inline_shapes))
print("paragraphs=", len(d.paragraphs))
print("two_column=", 'w:num="2"' in xml)
print("document_xml_bytes=", len(xml))
print("non_black_runs=", colored_runs)
print("underlined_runs=", underlined_runs)
expected_emails = {
    "divyaambre21@gmail.com",
    "sumitghogare2007@gmail.com",
    "tanvip1117@gmail.com",
    "tanavishrirame857@gmail.com",
}
blue_underlined = {
    run.text.strip()
    for paragraph in all_paragraphs
    for run in paragraph.runs
    if run.text.strip() in expected_emails
    and str(run.font.color.rgb) == "1155CC"
    and bool(run.font.underline)
}
title_style = d.styles["Title"]
heading_styles = [d.styles["Heading 1"], d.styles["Heading 2"]]
print("emails_blue_underlined=", blue_underlined == expected_emails)
print("title_black=", str(title_style.font.color.rgb) == "000000")
print("headings_black=", all(str(style.font.color.rgb) == "000000" for style in heading_styles))
print("title_border_absent=", title_style._element.get_or_add_pPr().find(qn("w:pBdr")) is None)
print("all_sections_continuous=", all(section.start_type == 0 for section in d.sections))
