from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "word" / "SmartFoodRescue_IEEE_Research_Paper.docx"
TMP = ROOT / "tmp" / "docx_assets"
OUT.parent.mkdir(parents=True, exist_ok=True)
TMP.mkdir(parents=True, exist_ok=True)


def font(size, bold=False):
    paths = [
        (r"C:\Windows\Fonts\timesbd.ttf" if bold else r"C:\Windows\Fonts\times.ttf"),
        (r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf"),
    ]
    for p in paths:
        if Path(p).exists():
            return ImageFont.truetype(p, size)
    return ImageFont.load_default()


def centered(draw, xy, text, fnt, fill="black"):
    x, y = xy
    box = draw.textbbox((0, 0), text, font=fnt)
    draw.text((x - (box[2]-box[0])/2, y - (box[3]-box[1])/2), text, font=fnt, fill=fill)


def arrow(draw, start, end, fill="#222222", width=5):
    draw.line([start, end], fill=fill, width=width)
    x2, y2 = end
    x1, y1 = start
    if abs(x2-x1) >= abs(y2-y1):
        s = 16 if x2 > x1 else -16
        draw.polygon([(x2, y2), (x2-s, y2-9), (x2-s, y2+9)], fill=fill)
    else:
        s = 16 if y2 > y1 else -16
        draw.polygon([(x2, y2), (x2-9, y2-s), (x2+9, y2-s)], fill=fill)


def diagram_architecture(path):
    im = Image.new("RGB", (1600, 660), "white")
    d = ImageDraw.Draw(im)
    fb = font(40, True)
    fs = font(30)
    role_y = 55
    for x, label in zip([60, 450, 840, 1230], ["Donor", "NGO", "Volunteer", "Administrator"]):
        d.rounded_rectangle((x, role_y, x+300, role_y+105), radius=16, fill="#DFF3E4", outline="#111111", width=5)
        centered(d, (x+150, role_y+54), label, fb)
        arrow(d, (x+150, role_y+105), (800, 235))
    d.rounded_rectangle((180, 235, 1420, 355), radius=18, fill="#E7F0FB", outline="#111111", width=5)
    centered(d, (800, 275), "React and TypeScript Client", fb)
    centered(d, (800, 325), "REST API   Socket.IO client   Browser geolocation", fs)
    arrow(d, (800, 355), (800, 425))
    d.rounded_rectangle((180, 425, 1140, 560), radius=18, fill="#FFF2CC", outline="#111111", width=5)
    centered(d, (660, 470), "Node.js Express and Socket.IO", fb)
    centered(d, (660, 525), "JWT RBAC workflow OAuth email and routing", fs)
    d.rounded_rectangle((1240, 425, 1515, 560), radius=18, fill="#F2F2F2", outline="#111111", width=5)
    centered(d, (1378, 470), "MongoDB Atlas", fb)
    centered(d, (1378, 525), "Persistent records", fs)
    arrow(d, (1140, 492), (1240, 492))
    im.save(path, dpi=(300, 300))


def diagram_lifecycle(path):
    im = Image.new("RGB", (1600, 560), "white")
    d = ImageDraw.Draw(im)
    fb = font(34, True)
    fi = font(29)
    top = [(45, "ASSIGNED", "#E7F0FB"), (430, "RECEIVED", "#E7F0FB"), (815, "DISPATCHED", "#E7F0FB"), (1200, "EN ROUTE", "#DFF3E4")]
    for x, label, fill in top:
        d.rounded_rectangle((x, 60, x+310, 155), radius=14, fill=fill, outline="#111111", width=5)
        centered(d, (x+155, 107), label, fb)
    for x in [355, 740, 1125]:
        arrow(d, (x, 107), (x+70, 107))
    arrow(d, (1355, 155), (1355, 255))
    bottom = [(430, "DISTRIBUTED"), (815, "DELIVERED"), (1200, "ARRIVED")]
    for x, label in bottom:
        d.rounded_rectangle((x, 255, x+310, 350), radius=14, fill="#DFF3E4", outline="#111111", width=5)
        centered(d, (x+155, 302), label, fb)
    arrow(d, (1200, 302), (1125, 302))
    arrow(d, (815, 302), (740, 302))
    centered(d, (800, 455), "Every transition is role checked and appended to the pickup audit trail", fi)
    im.save(path, dpi=(300, 300))


ARCH = TMP / "architecture.png"
LIFE = TMP / "lifecycle.png"
diagram_architecture(ARCH)
diagram_lifecycle(LIFE)


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.2677)
sec.page_height = Inches(11.6929)
sec.top_margin = Inches(0.55)
sec.bottom_margin = Inches(0.55)
sec.left_margin = Inches(0.63)
sec.right_margin = Inches(0.63)

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
normal.font.size = Pt(9)
normal.paragraph_format.line_spacing = 1.0
normal.paragraph_format.space_after = Pt(3)
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

title_style = doc.styles["Title"]
title_style.font.name = "Times New Roman"
title_style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
title_style.font.size = Pt(17)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)
title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_style.paragraph_format.space_after = Pt(8)
title_rpr = title_style._element.get_or_add_rPr()
title_color = title_rpr.find(qn("w:color"))
if title_color is not None:
    title_color.set(qn("w:val"), "000000")
    for attr in ("themeColor", "themeTint", "themeShade"):
        title_color.attrib.pop(qn(f"w:{attr}"), None)
title_ppr = title_style._element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

for style_name in ["Heading 1", "Heading 2"]:
    st = doc.styles[style_name]
    st.font.name = "Times New Roman"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.font.bold = True
    st.font.size = Pt(10)
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    st.paragraph_format.space_before = Pt(6)
    st.paragraph_format.space_after = Pt(3)
    st.paragraph_format.keep_with_next = True
    style_color = st._element.get_or_add_rPr().find(qn("w:color"))
    if style_color is not None:
        style_color.set(qn("w:val"), "000000")
        for attr in ("themeColor", "themeTint", "themeShade"):
            style_color.attrib.pop(qn(f"w:{attr}"), None)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)
        tc_borders.append(el)


def set_cell_margin(cell, top=80, start=90, bottom=80, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for tag, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        mar.append(node)


def format_table(table, widths):
    table.autofit = False
    for r_idx, row in enumerate(table.rows):
        if r_idx == 0:
            set_repeat_table_header(row)
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for c_idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[c_idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_borders(cell)
            set_cell_margin(cell)
            shade(cell, "D9E2F3" if r_idx == 0 else ("F2F2F2" if r_idx % 2 == 0 else "FFFFFF"))
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
                    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
                    run.font.size = Pt(7.5)
                    if r_idx == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(0, 0, 0)


def add_heading(text):
    return doc.add_paragraph(text, style="Heading 1")


def add_body(text, first_indent=True):
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(0.14) if first_indent else Inches(0)
    p.paragraph_format.keep_together = False
    p.add_run(text)
    return p


def add_caption(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = False
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(8)
    return p


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_end])


title = doc.add_paragraph(style="Title")
title_run = title.add_run("SmartFoodRescue Secure Full Stack Platform for\nTraceable Surplus Food Redistribution and Live Delivery Tracking")
title_run.font.color.rgb = RGBColor(0, 0, 0)
title_run.font.underline = False
title_ppr = title._p.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)

author_table = doc.add_table(rows=2, cols=3)
author_table.autofit = False
author_data = [
    ("1st Divya Ambre", "divyaambre21@gmail.com"),
    ("2nd Sumit Ghogare", "sumitghogare2007@gmail.com"),
    ("3rd Tanvi Patil", "tanvip1117@gmail.com"),
]
for idx, (name, email) in enumerate(author_data):
    cell = author_table.cell(0, idx)
    cell.width = Inches(2.33)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    lines = [name, "Department Of CSE(AILM)", "Vishwkarma Institute Of Technology", "Pune India"]
    for line_idx, line in enumerate(lines):
        r = p.add_run(line)
        r.font.name = "Times New Roman"
        r.font.size = Pt(8.3 if line_idx == 0 else 7.8)
        if line_idx < len(lines) - 1:
            r.add_break()
    r = p.add_run("\n" + email)
    r.font.name = "Times New Roman"
    r.font.size = Pt(7.8)
    r.font.color.rgb = RGBColor(17, 85, 204)
    r.font.underline = True

fourth = author_table.cell(1, 0).merge(author_table.cell(1, 2))
fourth.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
p = fourth.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(3)
p.paragraph_format.space_after = Pt(0)
for line_idx, line in enumerate(["4th Tanvi Shrirame", "Department Of CSE(AILM)", "Vishwkarma Institute Of Technology", "Pune India"]):
    r = p.add_run(line)
    r.font.name = "Times New Roman"
    r.font.size = Pt(8.3 if line_idx == 0 else 7.8)
    if line_idx < 3:
        r.add_break()
r = p.add_run("\n" + "tanavishrirame857@gmail.com")
r.font.name = "Times New Roman"
r.font.size = Pt(7.8)
r.font.color.rgb = RGBColor(17, 85, 204)
r.font.underline = True

for row in author_table.rows:
    for cell in row.cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_borders = OxmlElement("w:tcBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "nil")
            tc_borders.append(el)
        tc_pr.append(tc_borders)

body_sec = doc.add_section(WD_SECTION.CONTINUOUS)
doc.sections[0].start_type = WD_SECTION.CONTINUOUS
section_break_paragraph = doc.paragraphs[-1]
section_break_paragraph.paragraph_format.space_before = Pt(0)
section_break_paragraph.paragraph_format.space_after = Pt(0)
section_break_paragraph.paragraph_format.line_spacing = Pt(1)
body_sec.top_margin = Inches(0.55)
body_sec.bottom_margin = Inches(0.55)
body_sec.left_margin = Inches(0.63)
body_sec.right_margin = Inches(0.63)
cols = body_sec._sectPr.xpath("./w:cols")[0]
cols.set(qn("w:num"), "2")
cols.set(qn("w:space"), "360")

for section in doc.sections:
    add_page_number(section.footer.paragraphs[0])

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(0)
p.paragraph_format.space_after = Pt(3)
p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
r = p.add_run("Abstract  ")
r.bold = True
r.italic = True
p.add_run("Food rescue depends on rapid coordination because edible surplus is perishable, stakeholders are distributed, and responsibility changes repeatedly between donation and final distribution. This paper presents SmartFoodRescue, a full stack web platform connecting donors, non governmental organizations, volunteers, and administrators in one traceable workflow. The system uses a React and TypeScript client, Node.js and Express services, MongoDB Atlas, JSON Web Token authentication, and role based authorization. A validated state machine records the pickup lifecycle from ASSIGNED through RECEIVED, DISPATCHED, EN ROUTE, ARRIVED, DELIVERED, and DISTRIBUTED. During transport, the volunteer browser acquires device coordinates and streams validated updates through Socket.IO to an authorization protected room dedicated to that pickup. Gmail API OAuth 2.0 supports event notifications and secure password reset delivery. Client and server production builds passed, while implementation inspection confirmed the workflow, security, and tracking mechanisms. The platform joins donation management, accountable last mile logistics, real time tracking, and final distribution in one deployment oriented architecture.")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(4)
r = p.add_run("Keywords  ")
r.bold = True
r.italic = True
p.add_run("food rescue  surplus food redistribution  real time GPS tracking  Socket.IO  MERN  role based access control  logistics traceability")

add_heading("I INTRODUCTION")
add_body("Food waste and food insecurity coexist, while usable food frequently has only a short interval in which it can be collected and consumed. Prior research describes digital platforms as coordination intermediaries connecting surplus holders with charities or consumers [1], [2]. Studies of food rescue operations also identify missing facilitator organizations, weak logistics, and inconsistent communication as causes of avoidable loss [3]. A useful rescue platform must therefore manage more than a donation listing: it must coordinate acceptance, assignment, physical movement, handover, and evidence of distribution.")
add_body("Existing systems commonly address selected parts of this chain. Food sharing applications emphasize discovery and exchange; allocation research emphasizes fair matching; and volunteer crowdsourcing research examines task notification [4], [5]. A recent IEEE food donation application uses the MERN stack, role specific dashboards, geographic information, and real time coordination [6]. SmartFoodRescue extends this direction by integrating a seven stage pickup state machine with room scoped live GPS, auditable state changes, OAuth based email events, secure password recovery, and separate permissions for donors, NGOs, volunteers, and administrators.")
add_body("The contribution is threefold: an end to end workflow persisting the transition from donation to beneficiary distribution; a protected real time telemetry path bound to a pickup and its authorized participants; and a deployment ready security design combining JWT authentication, role aware REST and Socket.IO authorization, hashed reset tokens, validated coordinates, and environment isolated credentials.")

add_heading("II RELATED WORK")
add_body("Ciulli et al. characterize digital food recovery platforms as circularity brokers that connect supply and demand actors and combine platforms with existing operational processes [1]. Principato et al. show that multi sided surplus platforms must balance digital capability with social, environmental, and economic objectives [2]. Aloysius and Ananda report coordination and logistics gaps in a developing country food rescue system and recommend web and mobile platforms with more structured pickup processes [3].")
add_body("Mertzanidis et al. automate assignment of rejected food loads to food banks while considering fairness and driver efficiency [4]. Manshadi and Rodilitz formulate volunteer notification as an online decision problem using Food Rescue U.S. data [5]. The closest application comparator is a 2026 IEEE paper describing a MERN platform linking donors, volunteers, and recipients with authentication, geographic information, and real time coordination [6]. Other work includes a mobile surplus food distribution application with pickup scheduling and geolocation [7] and the AI FEED stakeholder informed food charity platform [8]. SmartFoodRescue differs by making physical delivery a first class auditable object with status history, an authorized socket room, current location, optional breadcrumbs, route information, and notifications.")

add_heading("III SYSTEM REQUIREMENTS AND ARCHITECTURE")
add_body("Four roles define the trust boundary. Donors create surplus food records with item name, category, quantity, expiry details, and pickup location. Authorized NGOs review available records and accept suitable donations. Volunteers receive assigned pickup tasks and progress them through collection and delivery. Administrators supervise users, activity, and active logistics. REST endpoints support normal data operations, while Socket.IO carries low latency location events.")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.keep_with_next = True
p.add_run().add_picture(str(ARCH), width=Inches(3.18))
add_caption("Fig 1  SmartFoodRescue logical architecture")
add_body("MongoDB stores users, role profiles, locations, donations, food items, requests, pickups, status history, distributions, current locations, and optional location history. The frontend can be deployed on Vercel, the backend on Render, and the database on MongoDB Atlas. Google Routes supplies road geometry and ETA when configured, while OSRM provides a routing fallback. Secrets and allowed origins are supplied through server environment variables.")

add_caption("TABLE I  IMPLEMENTATION STACK")
data = [
    ["Layer", "Technology", "Responsibility"],
    ["Client", "React TypeScript", "Role dashboards maps and device GPS"],
    ["API", "Node.js Express", "CRUD validation JWT and RBAC"],
    ["Realtime", "Socket.IO", "Pickup rooms and location events"],
    ["Persistence", "MongoDB Mongoose", "Workflow audit and tracking records"],
    ["Messaging", "Gmail API OAuth 2.0", "Lifecycle and reset emails"],
]
tbl = doc.add_table(rows=len(data), cols=3)
for i, row in enumerate(data):
    for j, value in enumerate(row):
        tbl.cell(i, j).text = value
format_table(tbl, [0.65, 0.95, 1.65])

add_heading("IV WORKFLOW AND DATA MODEL")
add_body("A donor submission creates a persistent donation that an NGO may inspect and accept according to role permissions. Acceptance establishes the logistics context and a pickup can be associated with a volunteer. The server rejects state jumps outside the transition graph shown in Fig 2. Every accepted transition creates a timestamped pickup tracking record that preserves who changed the state and when.")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.keep_with_next = True
p.add_run().add_picture(str(LIFE), width=Inches(3.18))
add_caption("Fig 2  Enforced pickup lifecycle")
add_body("The lifecycle separates operational events that simple applications often collapse. RECEIVED records that the volunteer took custody at the donor; DISPATCHED indicates departure readiness; EN ROUTE activates navigation; ARRIVED ends location transmission; DELIVERED records NGO handover; and DISTRIBUTED closes the loop after beneficiary distribution. Donation, pickup, tracking, and distribution records are synchronized while retaining immutable history.")
add_body("Current location is separated from optional location history. The current record stores latitude, longitude, accuracy, speed, heading, status, and update time for efficient reads. Breadcrumb records can support route replay or audit, subject to an explicit privacy and retention policy. Geographic fields are bounded and validated before persistence or broadcast.")

add_heading("V REAL TIME TRACKING AND NAVIGATION")
add_body("When the volunteer selects Start Navigation, the server validates that the pickup is DISPATCHED and advances it to EN ROUTE. The mobile browser calls navigator.geolocation.watchPosition(). Each observation is sent over an authenticated socket and associated with the pickup identifier. The backend verifies the user, pickup relationship, and coordinate bounds before updating the current location and broadcasting only to the pickup room.")
add_body("The join handler permits access only when the authenticated user is the assigned volunteer, owning donor, accepting NGO, or authorized administrator. Unrelated accounts cannot subscribe to a pickup movement. The interface displays origin, destination, volunteer marker, route, remaining distance, ETA, GPS accuracy, connection state, and last update time. Selecting Arrived requires EN ROUTE, preserves the last coordinate, and stops the browser watcher.")

add_heading("VI SECURITY AND NOTIFICATION DESIGN")
add_body("Passwords are hashed through the user model, while successful authentication issues a signed JWT used by HTTP and Socket.IO middleware. Route guards enforce role permissions. Sensitive donor and volunteer fields are returned only where operationally necessary. Helmet, server side validation, explicit cross origin configuration, and environment variables reduce deployment exposure.")
add_body("Gmail API OAuth 2.0 sends from the configured SmartFoodRescue mailbox without embedding a Gmail password. Event templates notify the donor when a request is submitted, accepted, assigned, delivered, or distributed. Recipient addresses come from the persisted user record rather than the client request.")
add_body("Password recovery generates 32 random bytes, sends the raw token only inside the reset URL, and stores its SHA 256 digest with a 30 minute expiry. The reset endpoint hashes the presented token and requires a matching unexpired digest. After a successful update, both token fields are cleared to provide single use behavior.")

add_caption("TABLE II  SECURITY CONTROLS")
data = [
    ["Threat", "Control"],
    ["Unauthorized API use", "JWT validation and role guards"],
    ["Tracking eavesdropping", "Authenticated pickup specific Socket.IO rooms"],
    ["Forged GPS payload", "Role assignment numeric range and accuracy validation"],
    ["Reset token database leak", "Random raw token with only SHA 256 digest persisted"],
    ["Credential disclosure", "Server environment variables and OAuth 2.0"],
]
tbl = doc.add_table(rows=len(data), cols=2)
for i, row in enumerate(data):
    for j, value in enumerate(row):
        tbl.cell(i, j).text = value
format_table(tbl, [1.05, 2.2])

add_heading("VII IMPLEMENTATION VERIFICATION AND DISCUSSION")
add_body("Verification distinguishes implemented behavior from future empirical evaluation. The TypeScript backend completed its production compilation. The React client completed the TypeScript and Vite production build, transforming 1,939 modules and producing deployable assets; Vite reported a non fatal chunk size warning. Repository inspection confirmed the seven state transition map, watchPosition(), authenticated pickup room handlers, room scoped broadcasts, OAuth Gmail integration, Google Routes with OSRM fallback, and hashed expiring reset token fields.")

add_caption("TABLE III  BUILD AND STATIC VERIFICATION")
data = [
    ["Verification item", "Observed result"],
    ["Backend production compile", "Passed"],
    ["Frontend TypeScript and Vite build", "Passed with optimization warning only"],
    ["Pickup transition graph", "Exact seven stage sequence enforced"],
    ["Realtime authorization", "Identity and pickup relationship checked before room join"],
    ["Location broadcast scope", "Pickup specific room only"],
    ["Reset lifecycle", "Random token stored hash expiry and post use clearing"],
]
tbl = doc.add_table(rows=len(data), cols=2)
for i, row in enumerate(data):
    for j, value in enumerate(row):
        tbl.cell(i, j).text = value
format_table(tbl, [1.35, 1.9])
add_body("These checks establish implementation completeness and compile time consistency, but they do not measure operational impact. Claims such as kilograms rescued, delivery time reduction, concurrent socket capacity, route accuracy, notification success, or beneficiary outcomes require a controlled pilot. The client bundle should also be code split before large scale deployment.")

add_heading("VIII LIMITATIONS AND FUTURE WORK")
add_body("The system depends on mobile permission, network coverage, and device GPS quality. Location histories introduce privacy risk and require consent, minimal retention, access logging, and deletion policies. OSRM route quality can vary, while commercial routing services introduce cost and quota constraints. Email delivery depends on OAuth configuration and provider limits. Food safety still requires organizational procedures beyond software, including temperature control, packaging, expiry interpretation, and liability guidance.")
add_body("Future work should evaluate the platform through a multi stakeholder pilot. Measures should include acceptance latency, pickup success rate, time in each state, expired donation rate, volunteer travel distance, route deviation, socket delay, notification success, and usability by role. Further additions may include proof of delivery media, offline GPS buffering, anomaly detection, multilingual interfaces, and privacy preserving impact dashboards.")

add_heading("IX CONCLUSION")
add_body("SmartFoodRescue integrates the social and technical stages of surplus food redistribution in one secure workflow. Its distinguishing feature is accountable custody and visibility from NGO acceptance through volunteer navigation and beneficiary distribution. The React, Express, MongoDB, JWT, Socket.IO, Gmail OAuth, and dual provider routing architecture is deployable with common cloud services and has passed client and server production builds. Separating REST transactions from protected realtime telemetry and enforcing a persistent pickup state machine provides a practical foundation for traceable, privacy aware food rescue operations. A field pilot is the necessary next step for quantifying impact and scalability.")

add_heading("X REFERENCES")
refs = [
    "[1] F. Ciulli, A. Kolk, and S. Boe Lillegraven, Circularity Brokers Digital Platform Organizations and Waste Recovery in Food Supply Chains, Journal of Business Ethics, vol. 167, pp. 299 to 331, 2020. https://doi.org/10.1007/s10551-019-04160-5",
    "[2] L. Principato et al., The influence of sustainability and digitalisation on business model innovation The case of a multi sided platform for food surplus redistribution, Industrial Marketing Management, vol. 115, pp. 156 to 171, 2023. https://doi.org/10.1016/j.indmarman.2023.09.001",
    "[3] N. Aloysius and J. Ananda, A Circular Economy Approach to Food Security and Poverty a Case Study in Food Rescue in Sri Lanka, Circular Economy and Sustainability, vol. 3, pp. 1919 to 1940, 2023. https://doi.org/10.1007/s43615-023-00255-4",
    "[4] M. Mertzanidis, A. Psomas, and P. Verma, Automating Food Drop The Power of Two Choices for Dynamic and Fair Food Allocation, arXiv 2406.06363, 2024. https://arxiv.org/abs/2406.06363",
    "[5] V. Manshadi and S. Rodilitz, Online Policies for Efficient Volunteer Crowdsourcing, arXiv 2002.08474, 2020. https://arxiv.org/abs/2002.08474",
    "[6] Bridging the Gap Between Donors and Recipients A Digital Application to Food Donation, Proceedings of ICCCES, 2026. https://doi.org/10.1109/ICCCES62661.2026.11436954",
    "[7] Y. A. Kusumawati et al., From Waste to Worth Enhancing Access to Edible Surplus Food Using a Mobile App for Better Distribution, Procedia Computer Science, vol. 269, pp. 749 to 761, 2025. https://doi.org/10.1016/j.procs.2025.09.018",
    "[8] M. Sammer et al., AI FEED Prototyping an AI Powered Platform for the Food Charity Ecosystem, International Journal of Computational Intelligence Systems, 2024. https://doi.org/10.1007/s44196-024-00656-9",
    "[9] M. Yu et al., Unlocking the potential of surplus food A blockchain approach to enhance equitable distribution and address food insecurity in Italy, Socio Economic Planning Sciences, vol. 93, art. 101868, 2024. https://doi.org/10.1016/j.seps.2024.101868",
    "[10] Digital platforms mapping the territory of new technologies to fight food waste, British Food Journal, vol. 122, no. 5, pp. 1647 to 1669, 2020. https://doi.org/10.1108/BFJ-06-2019-0391",
]
for ref in refs:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(ref)
    r.font.size = Pt(7.5)

doc.core_properties.title = "SmartFoodRescue Secure Full Stack Platform for Traceable Surplus Food Redistribution and Live Delivery Tracking"
doc.core_properties.subject = "IEEE style research paper on a secure food rescue and realtime delivery tracking platform"
doc.core_properties.author = "Divya Ambre; Sumit Ghogare; Tanvi Patil; Tanvi Shrirame"
doc.save(OUT)
print(OUT)
